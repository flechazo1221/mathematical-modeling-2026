"""Small, conservative finite-volume prototypes for the H1 M1--M4 candidates.

This is deliberately a prototype, not the full COMPUTE solver.  It reads the
original workbooks without modifying them and writes only below 03-prototype.
"""

from __future__ import annotations

import csv
import json
import math
import platform
import sys
import time
from dataclasses import dataclass
from pathlib import Path

import numpy as np
import openpyxl


ROOT = Path(__file__).resolve().parents[1]
OUT = Path(__file__).resolve().parent
SEED = 20260911
RADIUS = 0.02
LENGTH = 0.25
SENSITIVITY = {
    "low_drying": {"h_mult": 0.8, "hm_mult": 0.8, "d_prefactor_mult": 0.8, "d_exponent_mult": 1.2},
    "high_drying": {"h_mult": 1.2, "hm_mult": 1.2, "d_prefactor_mult": 1.2, "d_exponent_mult": 0.8},
}


def locate_xlsx(size: int) -> Path:
    matches = [p for p in (ROOT / "input").rglob("*.xlsx") if p.stat().st_size == size]
    if len(matches) != 1:
        raise RuntimeError(f"expected one xlsx of size {size}, found {len(matches)}")
    return matches[0]


def read_inputs():
    p1, p2 = locate_xlsx(16586), locate_xlsx(11485)
    wb1 = openpyxl.load_workbook(p1, read_only=True, data_only=True)
    ws1 = wb1[wb1.sheetnames[0]]
    rows1 = list(ws1.iter_rows(values_only=True))
    if (ws1.max_row, ws1.max_column) != (242, 3) or rows1[1][0] != 0 or rows1[-1][0] != 14400:
        raise RuntimeError("attachment 1 row/endpoint audit failed")
    env = np.asarray(rows1[1:], dtype=float)
    wb2 = openpyxl.load_workbook(p2, read_only=True, data_only=True)
    ws2 = wb2[wb2.sheetnames[0]]
    rows2 = list(ws2.iter_rows(values_only=True))
    if (ws2.max_row, ws2.max_column) != (146, 2) or rows2[1][0] != 0 or rows2[-1][0] != 259200:
        raise RuntimeError("attachment 2 row/endpoint audit failed")
    radius = np.asarray(rows2[1:], dtype=float)
    return p1, env, p2, radius


P1, ENV, P2, RADIUS_DATA = read_inputs()


def outside(t: float):
    tc = float(np.interp(min(t, ENV[-1, 0]), ENV[:, 0], ENV[:, 1]))
    c = float(np.interp(min(t, ENV[-1, 0]), ENV[:, 0], ENV[:, 2]))
    return tc + 273.15, c


@dataclass
class Grid:
    nr: int
    nz: int

    def __post_init__(self):
        self.rf = np.linspace(0.0, RADIUS, self.nr + 1)
        self.rc = 0.5 * (self.rf[:-1] + self.rf[1:])
        self.zf = np.linspace(0.0, LENGTH, self.nz + 1)
        self.zc = 0.5 * (self.zf[:-1] + self.zf[1:])
        self.dr = RADIUS / self.nr
        self.dz = LENGTH / self.nz
        self.vol = (math.pi * (self.rf[1:] ** 2 - self.rf[:-1] ** 2))[:, None] * self.dz
        self.annulus = math.pi * (self.rf[1:] ** 2 - self.rf[:-1] ** 2)


def harmonic(a, b):
    return 2.0 * a * b / np.maximum(a + b, 1e-300)


def step_diffusion(u, coeff, storage, grid: Grid, dt, conv, uenv, axial):
    net = np.zeros_like(u)
    # Radial internal faces.
    for i in range(grid.nr - 1):
        area = 2.0 * math.pi * grid.rf[i + 1] * grid.dz
        g = harmonic(coeff[i, :], coeff[i + 1, :]) * area / grid.dr
        q = g * (u[i + 1, :] - u[i, :])
        net[i, :] += q
        net[i + 1, :] -= q
    # Outer cylindrical Robin face; center face is symmetry/no flux.
    area = 2.0 * math.pi * RADIUS * grid.dz
    g = area / (0.5 * grid.dr / coeff[-1, :] + 1.0 / conv)
    q = g * (uenv - u[-1, :])
    net[-1, :] += q
    boundary_in = float(np.sum(q))
    if axial:
        for j in range(grid.nz - 1):
            area = grid.annulus
            g = harmonic(coeff[:, j], coeff[:, j + 1]) * area / grid.dz
            q = g * (u[:, j + 1] - u[:, j])
            net[:, j] += q
            net[:, j + 1] -= q
        for j in (0, grid.nz - 1):
            area = grid.annulus
            g = area / (0.5 * grid.dz / coeff[:, j] + 1.0 / conv)
            q = g * (uenv - u[:, j])
            net[:, j] += q
            boundary_in += float(np.sum(q))
    updated = u + dt * net / (storage * grid.vol)
    return updated, boundary_in


def props_q2(c, temp_k):
    rho = 650.0 + 128.0 * c
    cp = 1450.0 + 2736.0 * c / (c + 1.0)
    k = 0.21 + 0.38 * c / (c + 1.0)
    d = 2.4e-3 * np.exp(-0.45 / np.maximum(c, 1e-9)) * np.exp(-3850.0 / temp_k)
    return rho, cp, k, d


def robin_surface(cell, coeff, half_width, conv, env):
    return (coeff * cell / half_width + conv * env) / (coeff / half_width + conv)


def simulate(model, duration, nr, nz, dt, sample_every=300.0, *, h_mult=1.0, hm_mult=1.0,
             d_prefactor_mult=1.0, d_exponent_mult=1.0):
    grid = Grid(nr, nz)
    axial = nz > 1
    c = np.full((nr, nz), 2.55, dtype=float)
    temp = np.full((nr, nz), 301.15, dtype=float)
    inv0 = float(np.sum(c * grid.vol))
    boundary_integral = 0.0
    trajectory = []
    next_sample = 0.0
    start = time.perf_counter()
    nsteps = int(round(duration / dt))
    for step in range(nsteps + 1):
        t = min(step * dt, duration)
        tenv, cenv = outside(t)
        if t + 1e-9 >= next_sample or step == nsteps:
            j = nz // 2
            if model in ("M1", "M2_Q1"):
                ks = 0.36
                ds = 7e-9 * math.exp(-0.89 / max(float(c[-1, j]), 1e-9))
                ts = robin_surface(float(temp[-1, j]), ks, 0.5 * grid.dr, 25.0 * h_mult, tenv)
            elif model == "M2_Q2":
                _, _, _, d0 = props_q2(np.asarray([[2.55]]), np.asarray([[301.15]]))
                ds = float(d0[0, 0]); ts = tenv
            elif model == "M3":
                _, _, _, ds0 = props_q2(np.asarray([[c[-1, j]]]), np.asarray([[tenv]]))
                ds = float(ds0[0, 0]); ts = tenv
            else:
                _, _, ks0, ds0 = props_q2(np.asarray([[c[-1, j]]]), np.asarray([[temp[-1, j]]]))
                ks = float(ks0[0, 0]); ds = float(ds0[0, 0])
                ts = robin_surface(float(temp[-1, j]), ks, 0.5 * grid.dr, 25.0 * h_mult, tenv)
            if model in ("M2_Q2", "M3", "M4"):
                ds *= d_prefactor_mult * math.exp(-0.45 * (d_exponent_mult - 1.0) / max(float(c[-1, j]), 1e-9))
            cs = robin_surface(float(c[-1, j]), ds, 0.5 * grid.dr, 8e-7 * hm_mult, cenv)
            trajectory.append((t, float(temp[0, j] - 273.15), float(ts - 273.15),
                               float(c[0, j]), float(cs),
                               float(np.sum(c * grid.vol) / (np.sum(grid.vol) * grid.nz))))
            next_sample += sample_every
        if step == nsteps:
            break
        if model in ("M1", "M2_Q1"):
            rho = np.full_like(c, 820.0); cp = np.full_like(c, 2600.0)
            k = np.full_like(c, 0.36)
            d = 7e-9 * np.exp(-0.89 / np.maximum(c, 1e-9))
            temp, _ = step_diffusion(temp, k, rho * cp, grid, dt, 25.0 * h_mult, tenv, axial)
        elif model == "M2_Q2":
            rho0, cp0, k0, d0 = props_q2(np.full_like(c, 2.55), np.full_like(c, 301.15))
            d = d0
            temp[:] = tenv
        elif model == "M3":
            temp[:] = tenv
            _, _, _, d = props_q2(c, np.full_like(c, tenv))
        elif model == "M4":
            rho, cp, k, d = props_q2(c, temp)
            temp, _ = step_diffusion(temp, k, rho * cp, grid, dt, 25.0 * h_mult, tenv, axial)
        else:
            raise ValueError(model)
        if model in ("M2_Q2", "M3", "M4"):
            d = d * d_prefactor_mult * np.exp(-0.45 * (d_exponent_mult - 1.0) / np.maximum(c, 1e-9))
        c, bflux = step_diffusion(c, d, np.ones_like(c), grid, dt, 8e-7 * hm_mult, cenv, axial)
        boundary_integral += bflux * dt
        if not np.all(np.isfinite(c)) or not np.all(np.isfinite(temp)):
            raise FloatingPointError(f"non-finite state in {model} at {t}")
    runtime = time.perf_counter() - start
    inv1 = float(np.sum(c * grid.vol))
    balance = abs((inv1 - inv0) - boundary_integral) / max(abs(inv1 - inv0), 1e-15)
    violations = []
    if float(c.min()) < 0: violations.append("negative moisture")
    if float(temp.min()) < 250 or float(temp.max()) > 400: violations.append("temperature outside prototype physical guard")
    result = {
        "model": model, "duration_s": duration, "nr": nr, "nz": nz, "dt_s": dt,
        "seed": SEED, "runtime_s": runtime, "converged": not violations,
        "constraint_violations": violations, "normalized_moisture_balance_error": balance,
        "final_center_C": trajectory[-1][3], "final_surface_C": trajectory[-1][4],
        "final_mean_C": float(np.sum(c * grid.vol) / (np.sum(grid.vol) * grid.nz)),
        "final_center_T_C": trajectory[-1][1],
        "final_surface_T_C": trajectory[-1][2],
        "min_C": float(c.min()), "max_C": float(c.max()),
        "h_mult": h_mult, "hm_mult": hm_mult, "d_prefactor_mult": d_prefactor_mult,
        "d_exponent_mult": d_exponent_mult,
        "command_contract": "python run.py",
    }
    return result, trajectory, c, temp


MODEL_CONFIG = {
    "M1": ("model-m1-1d", [("coarse", 6, 1, 4.0), ("base", 10, 1, 2.0), ("fine", 14, 1, 1.0)], 1800.0),
    "M2": ("model-m2-2d", [("coarse", 6, 8, 4.0), ("base", 10, 12, 2.0), ("fine", 14, 16, 1.0)], 1800.0),
    "M3": ("model-m3-nonlinear", [("base", 10, 12, 2.0), ("time_refined", 10, 12, 1.0)], 10800.0),
    "M4": ("model-m4-coupled", [("base", 10, 12, 2.0), ("time_refined", 10, 12, 1.0)], 10800.0),
}


def run_one(label):
    dirname, cases, duration = MODEL_CONFIG[label]
    dpath = OUT / dirname
    dpath.mkdir(exist_ok=True)
    results = []
    saved_fields = {}
    for case, nr, nz, dt in cases:
        internal = "M2_Q1" if label == "M2" else label
        res, traj, c, temp = simulate(internal, duration, nr, nz, dt)
        res["candidate"] = label; res["case"] = case
        results.append(res)
        if case == "base":
            with (dpath / "trajectory.csv").open("w", newline="", encoding="utf-8-sig") as f:
                w = csv.writer(f); w.writerow(["time_s","center_T_C","surface_T_C","center_C","surface_C","mean_C"]); w.writerows(traj)
            saved_fields = {"C": c.tolist(), "T_C": (temp - 273.15).tolist()}
    # M2 also supplies the constant-property Q2 baseline on the same slice as M3/M4.
    if label == "M2":
        res, traj, c, temp = simulate("M2_Q2", 10800.0, 10, 12, 2.0)
        res["candidate"] = label; res["case"] = "q2_constant_baseline"
        results.append(res)
    if label in ("M3", "M4"):
        for case, factors in SENSITIVITY.items():
            res, _, _, _ = simulate(label, 10800.0, 10, 12, 2.0, **factors)
            res["candidate"] = label; res["case"] = case
            results.append(res)
    payload = {
        "candidate": label,
        "status": "PROTOTYPE_RUN_COMPLETE",
        "scope_note": "controlled small-slice prototype; not full COMPUTE",
        "results": results,
        "base_final_fields": saved_fields,
        "environment": {"python": sys.version, "platform": platform.platform(), "numpy": np.__version__, "openpyxl": openpyxl.__version__},
        "inputs": [str(P1.relative_to(ROOT)), str(P2.relative_to(ROOT))],
    }
    (dpath / "result.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return payload


def run_all():
    all_payload = [run_one(x) for x in ("M1", "M2", "M3", "M4")]
    rows = []
    for p in all_payload:
        for r in p["results"]:
            rows.append(r)
    fields = ["candidate","model","case","duration_s","nr","nz","dt_s","seed","runtime_s","converged",
              "constraint_violations","normalized_moisture_balance_error","final_center_C","final_surface_C","final_mean_C",
              "final_center_T_C","final_surface_T_C","min_C","max_C","h_mult","hm_mult",
              "d_prefactor_mult","d_exponent_mult","command_contract"]
    with (OUT / "原型结果.csv").open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=fields); w.writeheader()
        for r in rows:
            q = dict(r); q["constraint_violations"] = ";".join(q["constraint_violations"]); w.writerow({k:q[k] for k in fields})
    return all_payload


if __name__ == "__main__":
    labels = sys.argv[1:] or ["all"]
    if labels == ["all"]: run_all()
    else:
        for label in labels: run_one(label.upper())

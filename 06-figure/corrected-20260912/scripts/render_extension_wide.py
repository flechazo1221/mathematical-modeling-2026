"""Render only FIG-DATA-EXTENSION-WIDE from its frozen source snapshot.

The script is deliberately dependency-light and does not recompute a model.
It preserves every source sample, the existing axes, units, and frozen tail
extension rule, while producing vector SVG/PDF and a 600 DPI PNG.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont, ImageOps
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen.canvas import Canvas


ROOT = Path(__file__).resolve().parents[3]
OUT = ROOT / "06-figure" / "corrected-20260912"
FIGURES = OUT / "figures"
PREVIEWS = OUT / "previews"
SNAPSHOTS = OUT / "data-snapshots"
CONTRACTS = OUT / "contracts"
QA_DIR = OUT / "qa"

FIGURE = "FIG-DATA-EXTENSION-WIDE"
SOURCE = ROOT / "input" / "A题" / "附件" / "附件1.xlsx"
SNAPSHOT = SNAPSHOTS / f"{FIGURE}.csv"
SNAPSHOT_META = SNAPSHOTS / f"{FIGURE}.meta.json"

SVG = FIGURES / f"{FIGURE}.svg"
PNG = FIGURES / f"{FIGURE}.png"
PDF = FIGURES / f"{FIGURE}.pdf"
COLOR_PREVIEW = PREVIEWS / f"{FIGURE}-color.png"
GRAY_PREVIEW = PREVIEWS / f"{FIGURE}-grayscale.png"
LEGACY_PREVIEW = PREVIEWS / f"{FIGURE}-preview.png"
QA_REPORT = QA_DIR / f"{FIGURE}-qa.json"

FONT_PATH = Path(r"C:\Windows\Fonts\simsun.ttc")
PDF_FONT_PATH = Path(r"C:\Windows\Fonts\STSONG.TTF")

# Existing figure contract: keep the same physical canvas and axis ranges.
WIDTH_MM, HEIGHT_MM = 340.0, 70.0
SVG_W, SVG_H = 3400.0, 700.0
PLOT_LEFT, PLOT_RIGHT = 150.0, 3340.0
PLOT_TOP, PLOT_BOTTOM = 78.0, 532.0
XMIN, XMAX = 0.0, 10.0
YMIN, YMAX = 25.0, 55.0
EXTENSION_H = 4.0

BLUE = "#0072B2"
ORANGE = "#D55E00"
BLACK = "#000000"
GRID = "#000000"

# Legend geometry is shared by SVG, PNG, and PDF so every export has the
# same horizontal alignment and vertically centered rows.
LEGEND_X, LEGEND_Y = 2440.0, 376.0
LEGEND_W, LEGEND_H = 850.0, 120.0
LEGEND_ROW_1 = LEGEND_Y + 36.0
LEGEND_ROW_2 = LEGEND_Y + 84.0
LEGEND_LINE_X1 = LEGEND_X + 38.0
LEGEND_LINE_X2 = LEGEND_X + 104.0
LEGEND_MARKER_X = LEGEND_X + 71.0
LEGEND_TEXT_X = LEGEND_X + 132.0

SNAPSHOT_COLUMNS = [
    "series",
    "source_kind",
    "source_row",
    "time_s",
    "time_h",
    "temperature_C",
]


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def number(value: float) -> str:
    """Stable compact decimal representation for snapshot text."""
    value = float(value)
    if value.is_integer():
        return str(int(value))
    return format(value, ".17g")


def read_source() -> list[tuple[float, float, int]]:
    workbook = load_workbook(SOURCE, data_only=True, read_only=True)
    sheet = workbook.active
    rows: list[tuple[float, float, int]] = []
    for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if len(row) < 2 or row[0] is None or row[1] is None:
            continue
        try:
            time_s = float(row[0])
            temperature = float(row[1])
        except (TypeError, ValueError):
            continue
        if not (math.isfinite(time_s) and math.isfinite(temperature)):
            continue
        rows.append((time_s, temperature, row_number))
    workbook.close()
    if len(rows) != 241:
        raise RuntimeError(f"附件1 numeric sample count changed: expected 241, got {len(rows)}")
    if any(rows[i][0] >= rows[i + 1][0] for i in range(len(rows) - 1)):
        raise RuntimeError("附件1 time samples are not strictly increasing")
    return rows


def write_snapshot_if_missing(source_rows: list[tuple[float, float, int]]) -> None:
    SNAPSHOTS.mkdir(parents=True, exist_ok=True)
    if SNAPSHOT.exists():
        return
    last_time_s, last_temperature, _ = source_rows[-1]
    rows: list[dict[str, str]] = []
    for time_s, temperature, source_row in source_rows:
        rows.append(
            {
                "series": "observed",
                "source_kind": "attachment1",
                "source_row": str(source_row),
                "time_s": number(time_s),
                "time_h": number(time_s / 3600.0),
                "temperature_C": number(temperature),
            }
        )
    for time_s in (last_time_s, last_time_s + EXTENSION_H * 3600.0):
        rows.append(
            {
                "series": "extension",
                "source_kind": "frozen_tail_extension",
                "source_row": "",
                "time_s": number(time_s),
                "time_h": number(time_s / 3600.0),
                "temperature_C": number(last_temperature),
            }
        )
    with SNAPSHOT.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=SNAPSHOT_COLUMNS, lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


def validate_snapshot(source_rows: list[tuple[float, float, int]]) -> dict:
    if not SNAPSHOT.exists():
        raise RuntimeError(f"missing frozen snapshot: {SNAPSHOT}")
    with SNAPSHOT.open("r", encoding="utf-8", newline="") as handle:
        reader = csv.DictReader(handle)
        if reader.fieldnames != SNAPSHOT_COLUMNS:
            raise RuntimeError(f"unexpected snapshot columns: {reader.fieldnames}")
        rows = list(reader)

    observed = [row for row in rows if row["series"] == "observed"]
    extension = [row for row in rows if row["series"] == "extension"]
    if len(observed) != len(source_rows) or len(extension) != 2:
        raise RuntimeError("frozen snapshot row count changed")
    for snap, source in zip(observed, source_rows):
        snap_time, snap_temp = float(snap["time_s"]), float(snap["temperature_C"])
        if not (math.isclose(snap_time, source[0], rel_tol=0.0, abs_tol=1e-12) and
                math.isclose(snap_temp, source[1], rel_tol=0.0, abs_tol=1e-12)):
            raise RuntimeError(f"source/snapshot mismatch at source row {source[2]}")

    last_time_s, last_temperature, _ = source_rows[-1]
    expected_extension = [last_time_s, last_time_s + EXTENSION_H * 3600.0]
    for row, expected_time in zip(extension, expected_extension):
        if not (math.isclose(float(row["time_s"]), expected_time, rel_tol=0.0, abs_tol=1e-12) and
                math.isclose(float(row["temperature_C"]), last_temperature, rel_tol=0.0, abs_tol=1e-12)):
            raise RuntimeError("frozen extension endpoint changed")

    return {
        "source_path": SOURCE.relative_to(ROOT).as_posix(),
        "source_sha256": sha256(SOURCE),
        "snapshot_path": SNAPSHOT.relative_to(ROOT).as_posix(),
        "snapshot_sha256": sha256(SNAPSHOT),
        "source_rows": len(source_rows),
        "snapshot_observed_rows": len(observed),
        "extension_rows": len(extension),
        "data_equal": True,
        "max_abs_diff": 0.0,
        "source_time_range_s": [source_rows[0][0], source_rows[-1][0]],
        "extension_time_range_h": [expected_extension[0] / 3600.0, expected_extension[1] / 3600.0],
        "extension_temperature_C": last_temperature,
    }


def write_snapshot_meta(checks: dict) -> None:
    expected = {
        "figure_id": FIGURE,
        "source": checks["source_path"],
        "source_sha256": checks["source_sha256"],
        "snapshot": checks["snapshot_path"],
        "snapshot_sha256": checks["snapshot_sha256"],
        "columns": SNAPSHOT_COLUMNS,
        "observed_rows": checks["source_rows"],
        "extension_rule": "4 h后采用末段冻结稳健值；本图保留现有a(end,2)口径",
        "xlim_h": [XMIN, XMAX],
        "ylim_C": [YMIN, YMAX],
        "data_claim_status": "unchanged",
    }
    if SNAPSHOT_META.exists():
        actual = json.loads(SNAPSHOT_META.read_text(encoding="utf-8"))
        for key in ("source_sha256", "snapshot_sha256", "columns", "observed_rows"):
            if actual.get(key) != expected[key]:
                raise RuntimeError(f"snapshot metadata changed: {key}")
        return
    SNAPSHOT_META.write_text(json.dumps(expected, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def sx(value: float) -> float:
    return PLOT_LEFT + (value - XMIN) / (XMAX - XMIN) * (PLOT_RIGHT - PLOT_LEFT)


def sy(value: float) -> float:
    return PLOT_BOTTOM - (value - YMIN) / (YMAX - YMIN) * (PLOT_BOTTOM - PLOT_TOP)


def svg_escape(value: str) -> str:
    return (value.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            .replace('"', "&quot;").replace("'", "&apos;"))


def svg_line(x1: float, y1: float, x2: float, y2: float, *, stroke: str = BLACK,
             width: float = 2.0, opacity: float | None = None) -> str:
    extra = f' opacity="{opacity}"' if opacity is not None else ""
    return f'<line x1="{x1:.3f}" y1="{y1:.3f}" x2="{x2:.3f}" y2="{y2:.3f}" stroke="{stroke}" stroke-width="{width}"{extra}/>'


def svg_text(text: str, x: float, y: float, *, size: float, anchor: str = "start",
             weight: str = "normal", transform: str = "") -> str:
    transform_attr = f' transform="{transform}"' if transform else ""
    return (f'<text x="{x:.3f}" y="{y:.3f}" text-anchor="{anchor}" '
            f'font-family="SimSun, STSong, serif" font-size="{size:.2f}" '
            f'font-weight="{weight}" fill="{BLACK}"{transform_attr}>{svg_escape(text)}</text>')


def svg_path(points: list[tuple[float, float]]) -> str:
    commands = [f"M {sx(points[0][0]):.3f} {sy(points[0][1]):.3f}"]
    commands.extend(f"L {sx(x):.3f} {sy(y):.3f}" for x, y in points[1:])
    return " ".join(commands)


def render_svg(source_rows: list[tuple[float, float, int]]) -> None:
    observed = [(time_s / 3600.0, temperature) for time_s, temperature, _ in source_rows]
    last_time_h, last_temperature = observed[-1]
    extension = [(last_time_h, last_temperature), (last_time_h + EXTENSION_H, last_temperature)]
    marker_indices = list(range(0, len(observed), 12))
    if marker_indices[-1] != len(observed) - 1:
        marker_indices.append(len(observed) - 1)

    parts = [
        f'<svg xmlns="http://www.w3.org/2000/svg" width="{WIDTH_MM:g}mm" height="{HEIGHT_MM:g}mm" viewBox="0 0 {SVG_W:g} {SVG_H:g}">',
        '<rect width="100%" height="100%" fill="#FFFFFF"/>',
        svg_text("末段稳健统计量延拓示意", SVG_W / 2, 43, size=31, anchor="middle", weight="bold"),
    ]
    for tick in range(0, 11):
        x = sx(float(tick))
        parts.append(svg_line(x, PLOT_TOP, x, PLOT_BOTTOM, stroke=GRID, width=1.2, opacity=0.14))
    for tick in range(25, 56, 5):
        y = sy(float(tick))
        parts.append(svg_line(PLOT_LEFT, y, PLOT_RIGHT, y, stroke=GRID, width=1.2, opacity=0.14))

    parts.extend([
        svg_line(PLOT_LEFT, PLOT_BOTTOM, PLOT_RIGHT, PLOT_BOTTOM, width=2.5),
        svg_line(PLOT_LEFT, PLOT_TOP, PLOT_LEFT, PLOT_BOTTOM, width=2.5),
    ])
    for tick in range(0, 11):
        x = sx(float(tick))
        parts.append(svg_line(x, PLOT_BOTTOM, x, PLOT_BOTTOM + 13, width=2.0))
        parts.append(svg_text(str(tick), x, PLOT_BOTTOM + 40, size=25, anchor="middle"))
    for tick in range(25, 56, 5):
        y = sy(float(tick))
        parts.append(svg_line(PLOT_LEFT - 13, y, PLOT_LEFT, y, width=2.0))
        parts.append(svg_text(str(tick), PLOT_LEFT - 24, y + 9, size=25, anchor="end"))

    parts.append(f'<path d="{svg_path(observed)}" fill="none" stroke="{BLUE}" stroke-width="5.4" stroke-linejoin="round" stroke-linecap="round"/>')
    for index in marker_indices:
        x, y = observed[index]
        parts.append(f'<circle cx="{sx(x):.3f}" cy="{sy(y):.3f}" r="7.0" fill="#FFFFFF" stroke="{BLUE}" stroke-width="3.0"/>')
    parts.append(f'<path d="{svg_path(extension)}" fill="none" stroke="{ORANGE}" stroke-width="5.4" stroke-linecap="round"/>')
    for x, y in extension:
        center_x, center_y = sx(x), sy(y)
        parts.append(f'<rect x="{center_x - 8:.3f}" y="{center_y - 8:.3f}" width="16" height="16" fill="#FFFFFF" stroke="{ORANGE}" stroke-width="3.0"/>')

    # Compact legend in the existing empty lower-right region; no extra notes.
    parts.append(f'<rect x="{LEGEND_X}" y="{LEGEND_Y}" width="{LEGEND_W}" height="{LEGEND_H}" fill="#FFFFFF" fill-opacity="0.94" stroke="{BLACK}" stroke-opacity="0.25" stroke-width="1.4"/>')
    parts.append(svg_line(LEGEND_LINE_X1, LEGEND_ROW_1, LEGEND_LINE_X2, LEGEND_ROW_1, stroke=BLUE, width=5.0))
    parts.append(f'<circle cx="{LEGEND_MARKER_X}" cy="{LEGEND_ROW_1}" r="6.5" fill="#FFFFFF" stroke="{BLUE}" stroke-width="2.7"/>')
    parts.append(svg_text("观测区间（分段线性）", LEGEND_TEXT_X, LEGEND_ROW_1 + 9.0, size=25))
    parts.append(svg_line(LEGEND_LINE_X1, LEGEND_ROW_2, LEGEND_LINE_X2, LEGEND_ROW_2, stroke=ORANGE, width=5.0))
    parts.append(f'<rect x="{LEGEND_MARKER_X - 8.0}" y="{LEGEND_ROW_2 - 8.0}" width="16" height="16" fill="#FFFFFF" stroke="{ORANGE}" stroke-width="2.7"/>')
    parts.append(svg_text("冻结延拓", LEGEND_TEXT_X, LEGEND_ROW_2 + 9.0, size=25))
    parts.append(svg_text("时间 (h)", (PLOT_LEFT + PLOT_RIGHT) / 2, 663, size=27, anchor="middle"))
    parts.append(svg_text("温度边界 (°C)", 43, (PLOT_TOP + PLOT_BOTTOM) / 2, size=27, anchor="middle", transform=f"rotate(-90 43 {(PLOT_TOP + PLOT_BOTTOM) / 2:.3f})"))
    parts.append("</svg>")
    SVG.write_text("\n".join(parts) + "\n", encoding="utf-8")


def font(size_px: int) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(str(FONT_PATH), size_px, index=0)


def pxy(x: float, y: float, width_px: int, height_px: int) -> tuple[int, int]:
    return (round(x / SVG_W * width_px), round(y / SVG_H * height_px))


def draw_png(source_rows: list[tuple[float, float, int]]) -> Image.Image:
    width_px = round(WIDTH_MM / 25.4 * 600)
    height_px = round(HEIGHT_MM / 25.4 * 600)
    image = Image.new("RGB", (width_px, height_px), "white")
    draw = ImageDraw.Draw(image)

    def line(points: list[tuple[float, float]], fill: str, width: int) -> None:
        draw.line([pxy(x, y, width_px, height_px) for x, y in points], fill=fill, width=width, joint="curve")

    def text(value: str, xy: tuple[float, float], size: int, anchor: str = "la", fill: str = BLACK) -> None:
        draw.text(pxy(xy[0], xy[1], width_px, height_px), value, font=font(size), fill=fill, anchor=anchor)

    def circle(x: float, y: float, radius: float, outline: str, width: int) -> None:
        cx, cy = pxy(x, y, width_px, height_px)
        rx = round(radius / SVG_W * width_px)
        ry = round(radius / SVG_H * height_px)
        draw.ellipse((cx - rx, cy - ry, cx + rx, cy + ry), fill="white", outline=outline, width=width)

    def square(x: float, y: float, half: float, outline: str, width: int) -> None:
        cx, cy = pxy(x, y, width_px, height_px)
        hx = round(half / SVG_W * width_px)
        hy = round(half / SVG_H * height_px)
        draw.rectangle((cx - hx, cy - hy, cx + hx, cy + hy), fill="white", outline=outline, width=width)

    text("末段稳健统计量延拓示意", (SVG_W / 2, 43), 73, "ma")
    for tick in range(0, 11):
        x = sx(float(tick))
        line([(x, PLOT_TOP), (x, PLOT_BOTTOM)], "#00000026", 2)
    for tick in range(25, 56, 5):
        y = sy(float(tick))
        line([(PLOT_LEFT, y), (PLOT_RIGHT, y)], "#00000026", 2)
    line([(PLOT_LEFT, PLOT_BOTTOM), (PLOT_RIGHT, PLOT_BOTTOM)], BLACK, 5)
    line([(PLOT_LEFT, PLOT_TOP), (PLOT_LEFT, PLOT_BOTTOM)], BLACK, 5)
    for tick in range(0, 11):
        x = sx(float(tick))
        line([(x, PLOT_BOTTOM), (x, PLOT_BOTTOM + 13)], BLACK, 4)
        text(str(tick), (x, PLOT_BOTTOM + 40), 59, "ma")
    for tick in range(25, 56, 5):
        y = sy(float(tick))
        line([(PLOT_LEFT - 13, y), (PLOT_LEFT, y)], BLACK, 4)
        text(str(tick), (PLOT_LEFT - 24, y + 1), 59, "rm")

    observed = [(time_s / 3600.0, temperature) for time_s, temperature, _ in source_rows]
    line([(sx(x), sy(y)) for x, y in observed], BLUE, 13)
    marker_indices = list(range(0, len(observed), 12))
    if marker_indices[-1] != len(observed) - 1:
        marker_indices.append(len(observed) - 1)
    for index in marker_indices:
        x, y = observed[index]
        circle(sx(x), sy(y), 7.0, BLUE, 7)
    last_time_h, last_temperature = observed[-1]
    extension = [(last_time_h, last_temperature), (last_time_h + EXTENSION_H, last_temperature)]
    line([(sx(x), sy(y)) for x, y in extension], ORANGE, 13)
    for x, y in extension:
        square(sx(x), sy(y), 8.0, ORANGE, 7)

    x0, y0 = pxy(LEGEND_X, LEGEND_Y, width_px, height_px)
    x1, y1 = pxy(LEGEND_X + LEGEND_W, LEGEND_Y + LEGEND_H, width_px, height_px)
    draw.rectangle((x0, y0, x1, y1), fill="#FFFFFF", outline="#00000040", width=3)
    line([(LEGEND_LINE_X1, LEGEND_ROW_1), (LEGEND_LINE_X2, LEGEND_ROW_1)], BLUE, 12)
    circle(LEGEND_MARKER_X, LEGEND_ROW_1, 6.5, BLUE, 6)
    text("观测区间（分段线性）", (LEGEND_TEXT_X, LEGEND_ROW_1), 59, "lm")
    line([(LEGEND_LINE_X1, LEGEND_ROW_2), (LEGEND_LINE_X2, LEGEND_ROW_2)], ORANGE, 12)
    square(LEGEND_MARKER_X, LEGEND_ROW_2, 8.0, ORANGE, 6)
    text("冻结延拓", (LEGEND_TEXT_X, LEGEND_ROW_2), 59, "lm")
    text("时间 (h)", ((PLOT_LEFT + PLOT_RIGHT) / 2, 663), 64, "ma")
    y_label = Image.new("RGBA", (height_px, 120), (255, 255, 255, 0))
    y_draw = ImageDraw.Draw(y_label)
    y_draw.text((height_px // 2, 60), "温度边界 (°C)", font=font(64), fill=BLACK, anchor="mm")
    y_label = y_label.rotate(90, expand=True, resample=Image.Resampling.BICUBIC)
    image.paste(y_label, (0, round(height_px / 2 - y_label.height / 2)), y_label)
    return image


def pdf_text(canvas: Canvas, value: str, x: float, y: float, size: float, anchor: str = "left") -> None:
    if anchor == "middle":
        canvas.drawCentredString(x, y, value)
    elif anchor == "right":
        canvas.drawRightString(x, y, value)
    else:
        canvas.drawString(x, y, value)


def render_pdf(source_rows: list[tuple[float, float, int]]) -> None:
    if "STSong" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("STSong", str(PDF_FONT_PATH)))
    page_w = WIDTH_MM / 25.4 * 72.0
    page_h = HEIGHT_MM / 25.4 * 72.0
    canvas = Canvas(str(PDF), pagesize=(page_w, page_h), pageCompression=1, invariant=1)
    canvas.setTitle(FIGURE)
    canvas.setCreator("Codex figure stage")
    canvas.setFillColorRGB(1, 1, 1)
    canvas.rect(0, 0, page_w, page_h, fill=1, stroke=0)

    def px(value: float) -> float:
        return value / SVG_W * page_w

    def py(value: float) -> float:
        return page_h - value / SVG_H * page_h

    def line(x1: float, y1: float, x2: float, y2: float, color: str = BLACK, width: float = 0.8) -> None:
        color = color.lstrip("#")
        canvas.setStrokeColorRGB(int(color[0:2], 16) / 255, int(color[2:4], 16) / 255, int(color[4:6], 16) / 255)
        canvas.setLineWidth(width)
        canvas.line(px(x1), py(y1), px(x2), py(y2))

    def text(value: str, x: float, y: float, size: float, anchor: str = "left") -> None:
        canvas.setFillColorRGB(0, 0, 0)
        canvas.setFont("STSong", size / SVG_W * page_w)
        pdf_text(canvas, value, px(x), py(y), size / SVG_W * page_w, anchor)

    text("末段稳健统计量延拓示意", SVG_W / 2, 43, 31, "middle")
    canvas.setStrokeColorRGB(0, 0, 0)
    canvas.setStrokeAlpha(0.14)
    for tick in range(0, 11):
        x = sx(float(tick))
        canvas.setLineWidth(0.45)
        canvas.line(px(x), py(PLOT_TOP), px(x), py(PLOT_BOTTOM))
    for tick in range(25, 56, 5):
        y = sy(float(tick))
        canvas.line(px(PLOT_LEFT), py(y), px(PLOT_RIGHT), py(y))
    canvas.setStrokeAlpha(1.0)
    line(PLOT_LEFT, PLOT_BOTTOM, PLOT_RIGHT, PLOT_BOTTOM, width=1.1)
    line(PLOT_LEFT, PLOT_TOP, PLOT_LEFT, PLOT_BOTTOM, width=1.1)
    for tick in range(0, 11):
        x = sx(float(tick))
        line(x, PLOT_BOTTOM, x, PLOT_BOTTOM + 13, width=0.9)
        text(str(tick), x, PLOT_BOTTOM + 40, 25, "middle")
    for tick in range(25, 56, 5):
        y = sy(float(tick))
        line(PLOT_LEFT - 13, y, PLOT_LEFT, y, width=0.9)
        text(str(tick), PLOT_LEFT - 24, y - 9, 25, "right")

    def rgb(hex_color: str) -> tuple[float, float, float]:
        value = hex_color.lstrip("#")
        return tuple(int(value[i:i + 2], 16) / 255 for i in (0, 2, 4))

    def curve(points: list[tuple[float, float]], color: str, width: float) -> None:
        canvas.setStrokeColorRGB(*rgb(color))
        canvas.setLineWidth(width)
        path = canvas.beginPath()
        path.moveTo(px(sx(points[0][0])), py(sy(points[0][1])))
        for x, y in points[1:]:
            path.lineTo(px(sx(x)), py(sy(y)))
        canvas.drawPath(path, stroke=1, fill=0)

    observed = [(time_s / 3600.0, temperature) for time_s, temperature, _ in source_rows]
    curve(observed, BLUE, 1.65)
    marker_indices = list(range(0, len(observed), 12))
    if marker_indices[-1] != len(observed) - 1:
        marker_indices.append(len(observed) - 1)
    for index in marker_indices:
        x, y = sx(observed[index][0]), sy(observed[index][1])
        canvas.setStrokeColorRGB(*rgb(BLUE))
        canvas.setFillColorRGB(1, 1, 1)
        canvas.setLineWidth(0.85)
        canvas.circle(px(x), py(y), px(7.0), fill=1, stroke=1)
    last_time_h, last_temperature = observed[-1]
    extension = [(last_time_h, last_temperature), (last_time_h + EXTENSION_H, last_temperature)]
    curve(extension, ORANGE, 1.65)
    canvas.setStrokeColorRGB(*rgb(ORANGE))
    canvas.setFillColorRGB(1, 1, 1)
    canvas.setLineWidth(0.85)
    for x, y in extension:
        cx, cy = px(sx(x)), py(sy(y))
        half = px(8.0)
        canvas.rect(cx - half, cy - half, half * 2, half * 2, fill=1, stroke=1)

    canvas.setFillColorRGB(1, 1, 1)
    canvas.setStrokeColorRGB(0, 0, 0)
    canvas.setStrokeAlpha(0.25)
    canvas.rect(px(LEGEND_X), py(LEGEND_Y + LEGEND_H), px(LEGEND_W), py(LEGEND_Y) - py(LEGEND_Y + LEGEND_H), fill=1, stroke=1)
    canvas.setStrokeAlpha(1.0)
    line(LEGEND_LINE_X1, LEGEND_ROW_1, LEGEND_LINE_X2, LEGEND_ROW_1, BLUE, 1.55)
    canvas.setStrokeColorRGB(*rgb(BLUE)); canvas.setFillColorRGB(1, 1, 1); canvas.setLineWidth(0.8)
    canvas.circle(px(LEGEND_MARKER_X), py(LEGEND_ROW_1), px(6.5), fill=1, stroke=1)
    text("观测区间（分段线性）", LEGEND_TEXT_X, LEGEND_ROW_1 + 8.0, 25)
    line(LEGEND_LINE_X1, LEGEND_ROW_2, LEGEND_LINE_X2, LEGEND_ROW_2, ORANGE, 1.55)
    canvas.setStrokeColorRGB(*rgb(ORANGE)); canvas.setFillColorRGB(1, 1, 1); canvas.setLineWidth(0.8)
    half = px(8.0); cx, cy = px(LEGEND_MARKER_X), py(LEGEND_ROW_2); canvas.rect(cx - half, cy - half, 2 * half, 2 * half, fill=1, stroke=1)
    text("冻结延拓", LEGEND_TEXT_X, LEGEND_ROW_2 + 8.0, 25)
    text("时间 (h)", (PLOT_LEFT + PLOT_RIGHT) / 2, 663, 27, "middle")
    canvas.saveState()
    canvas.translate(px(43), py((PLOT_TOP + PLOT_BOTTOM) / 2))
    canvas.rotate(90)
    canvas.setFillColorRGB(0, 0, 0)
    canvas.setFont("STSong", 27 / SVG_W * page_w)
    canvas.drawCentredString(0, 0, "温度边界 (°C)")
    canvas.restoreState()
    canvas.showPage()
    canvas.save()


def write_qa(checks: dict, image: Image.Image) -> None:
    QA_DIR.mkdir(parents=True, exist_ok=True)
    qa = {
        "figure_id": FIGURE,
        "data_check": checks,
        "visual_contract": {
            "data_curves": "高对比彩色；观测段圆 marker，冻结延拓段方 marker；两段均为实线",
            "non_data_elements": "坐标轴、网格、标题、图例、刻度和文字统一黑色",
            "font": "SimSun / 宋体",
            "xlim_h": [XMIN, XMAX],
            "ylim_C": [YMIN, YMAX],
            "note_policy": "删除图内非必要延拓起点文字和备注；限制条件只进图注/正文",
        },
        "outputs": {
            "png": PNG.relative_to(ROOT).as_posix(),
            "svg": SVG.relative_to(ROOT).as_posix(),
            "pdf": PDF.relative_to(ROOT).as_posix(),
            "color_preview": COLOR_PREVIEW.relative_to(ROOT).as_posix(),
            "grayscale_preview": GRAY_PREVIEW.relative_to(ROOT).as_posix(),
            "legacy_preview_alias": LEGACY_PREVIEW.relative_to(ROOT).as_posix(),
            "png_pixels": list(image.size),
            "png_dpi": [600, 600],
        },
    }
    for key, path in (("png_sha256", PNG), ("svg_sha256", SVG), ("pdf_sha256", PDF),
                      ("color_preview_sha256", COLOR_PREVIEW), ("grayscale_preview_sha256", GRAY_PREVIEW),
                      ("legacy_preview_sha256", LEGACY_PREVIEW)):
        qa["outputs"][key] = sha256(path)
    QA_REPORT.write_text(json.dumps(qa, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser(description="Render only FIG-DATA-EXTENSION-WIDE")
    parser.parse_args()
    if not SOURCE.exists() or not FONT_PATH.exists() or not PDF_FONT_PATH.exists():
        raise FileNotFoundError("source workbook or required Chinese font is missing")
    for directory in (FIGURES, PREVIEWS, SNAPSHOTS, CONTRACTS, QA_DIR):
        directory.mkdir(parents=True, exist_ok=True)
    source_rows = read_source()
    write_snapshot_if_missing(source_rows)
    checks = validate_snapshot(source_rows)
    write_snapshot_meta(checks)

    render_svg(source_rows)
    image = draw_png(source_rows)
    image.save(PNG, dpi=(600, 600))
    # A half-scale preview is still rendered from the formal 600 DPI bitmap.
    preview = image.resize((image.width // 2, image.height // 2), Image.Resampling.LANCZOS)
    preview.save(COLOR_PREVIEW, dpi=(300, 300))
    ImageOps.grayscale(preview).save(GRAY_PREVIEW, dpi=(300, 300))
    # Keep the pre-existing handoff path as an alias of the revised color preview.
    preview.save(LEGACY_PREVIEW, dpi=(300, 300))
    render_pdf(source_rows)
    write_qa(checks, image)
    print(json.dumps({
        "figure_id": FIGURE,
        "source_sha256": checks["source_sha256"],
        "snapshot_sha256": checks["snapshot_sha256"],
        "data_equal": checks["data_equal"],
        "png_pixels": list(image.size),
        "png_sha256": sha256(PNG),
        "svg_sha256": sha256(SVG),
        "pdf_sha256": sha256(PDF),
        "color_preview_sha256": sha256(COLOR_PREVIEW),
        "grayscale_preview_sha256": sha256(GRAY_PREVIEW),
        "legacy_preview_sha256": sha256(LEGACY_PREVIEW),
    }, ensure_ascii=False))


if __name__ == "__main__":
    main()

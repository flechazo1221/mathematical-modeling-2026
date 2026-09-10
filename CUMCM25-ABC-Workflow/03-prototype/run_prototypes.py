from __future__ import annotations

import csv, hashlib, json, math, os, platform, sys, time
from pathlib import Path
import numpy as np
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "03-prototype"
CFG = json.loads((OUT / "预注册规则.json").read_text(encoding="utf-8"))
SEED = int(CFG["seed"])
FILES = {
    ("SiC", 10): "附件1.xlsx", ("SiC", 15): "附件2.xlsx",
    ("Si", 10): "附件3.xlsx", ("Si", 15): "附件4.xlsx",
}

def read_xlsx(path: Path):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    vals = list(ws.iter_rows(min_row=2, values_only=True))
    sigma = np.asarray([float(x[0]) for x in vals])
    refl = np.asarray([float(x[1]) for x in vals])
    wb.close()
    return sigma, refl

def detrend(x, y, degree=2):
    z = 2 * (x - x.min()) / (x.max() - x.min()) - 1
    co = np.polynomial.polynomial.polyfit(z, y, degree)
    return y - np.polynomial.polynomial.polyval(z, co)

def estimate_B(x, y):
    yy = detrend(x, y)
    s = np.sign(yy)
    idx = np.where(s[:-1] * s[1:] < 0)[0]
    if len(idx) < 3: return np.nan, "INSUFFICIENT_FRINGES", np.nan
    roots = x[idx] - yy[idx] * (x[idx+1]-x[idx]) / (yy[idx+1]-yy[idx])
    half = np.diff(roots)
    half = half[(half > 0) & np.isfinite(half)]
    if len(half) < 2: return np.nan, "FRINGE_AMBIGUOUS", np.nan
    freq = 1.0 / (2.0 * np.median(half))
    q_um = freq * 1e4 / 2.0
    ambiguity = float(np.median(np.abs(half-np.median(half))) / np.median(half))
    return q_um, "OK_CONDITIONAL", ambiguity

def estimate_F(x, y):
    n = len(x); grid = np.linspace(x.min(), x.max(), n)
    yy = np.interp(grid, x, y); yy = detrend(grid, yy)
    win = np.hanning(n); spec = np.abs(np.fft.rfft(yy*win))
    freq = np.fft.rfftfreq(n, d=grid[1]-grid[0]); spec[:2] = 0
    k = int(np.argmax(spec))
    if k <= 1 or k >= len(spec)-1: return np.nan, "PEAK_AMBIGUOUS", np.nan
    local = spec[max(2,k-2):min(len(spec),k+3)]
    second = np.partition(spec[2:], -2)[-2]
    ambiguity = float(second/spec[k]) if spec[k] else 1.0
    q_um = freq[k]*1e4/2
    return q_um, "OK_CONDITIONAL", ambiguity

def fit_at_frequency(x, y, f):
    z = 2*(x-x.min())/(x.max()-x.min())-1
    A = np.column_stack([np.ones_like(x), z, np.cos(2*np.pi*f*x), np.sin(2*np.pi*f*x)])
    beta, *_ = np.linalg.lstsq(A, y, rcond=None)
    resid = y-A@beta
    rss = float(resid@resid)
    return rss, resid

def estimate_P(x, y):
    q0, st, _ = estimate_F(x,y)
    if not np.isfinite(q0): return np.nan, "MULTISTART_UNSTABLE", np.nan
    f0 = 2*q0/1e4
    lo=max(1/(x.max()-x.min()), f0*0.7); hi=f0*1.3
    freqs=np.linspace(lo,hi,401)
    rss=np.asarray([fit_at_frequency(x,y,f)[0] for f in freqs])
    k=int(np.argmin(rss)); f=freqs[k]
    if k in (0,len(freqs)-1): return np.nan,"BOUNDARY_SOLUTION",np.nan
    curvature=(rss[k-1]+rss[k+1]-2*rss[k])/max(rss[k],1e-30)
    return f*1e4/2,"OK_CONDITIONAL",float(curvature)

EST={"B":estimate_B,"F":estimate_F,"P":estimate_P}

def synthetic_rows():
    rng=np.random.default_rng(SEED); rows=[]
    base_sigma,_=read_xlsx(ROOT/"input"/"B题"/"附件"/"附件1.xlsx")
    for wi,(a,b) in enumerate(CFG["comparison"]["windows_cm_inv"],1):
      x=base_sigma[(base_sigma>=a)&(base_sigma<=b)]
      for q in CFG["synthetic_calibration"]["optical_thickness_um"]:
       for noise in CFG["synthetic_calibration"]["noise_sd_fraction"]:
        for slope in CFG["synthetic_calibration"]["baseline_slope_fraction"]:
         for rep in range(CFG["synthetic_calibration"]["replicates"]):
          phase=rng.uniform(0,2*np.pi); amp=.25
          y=.5+amp*np.cos(4*np.pi*q*1e-4*x+phase)+slope*(x-x.mean())/(x.max()-x.min())+rng.normal(0,noise*amp,len(x))
          cycles=2*q*1e-4*(b-a)
          for model,fn in EST.items():
           est,st,diag=fn(x,y); err=abs(est-q)/q if np.isfinite(est) else np.nan
           lim=max(.02,2/max(cycles,1e-12))
           rows.append(dict(kind="synthetic",model=model,material="CAL",angle_deg="",window=f"W{wi}",mask="na",q_um=est,status=st,diagnostic=diag,relative_error=err,threshold=lim,pass_recovery=bool(np.isfinite(err) and err<=lim),runtime_s=""))
    return rows

def observed_rows():
    rows=[]
    for (mat,angle),name in FILES.items():
      sigma,refl=read_xlsx(ROOT/"input"/"B题"/"附件"/name)
      reasoned=np.isfinite(sigma)&np.isfinite(refl)&(refl<=100)
      reasoned[0]=False
      for wi,(a,b) in enumerate(CFG["comparison"]["windows_cm_inv"],1):
       base=(sigma>=a)&(sigma<=b)
       for mask_name,mask in [("raw",base),("reasoned",base&reasoned)]:
        x=sigma[mask]; y=refl[mask]
        for model,fn in EST.items():
         t=time.perf_counter(); est,st,diag=fn(x,y); rt=time.perf_counter()-t
         rows.append(dict(kind="observed",model=model,material=mat,angle_deg=angle,window=f"W{wi}",mask=mask_name,q_um=est,status=st,diagnostic=diag,relative_error="",threshold="",pass_recovery="",runtime_s=rt))
    return rows

def summarize(rows):
    syn=[r for r in rows if r["kind"]=="synthetic"]
    obs=[r for r in rows if r["kind"]=="observed"]
    summary={}
    bmed=np.nanmedian([r["relative_error"] for r in syn if r["model"]=="B"])
    for m in "BFP":
      ss=[r for r in syn if r["model"]==m]
      errs=np.asarray([r["relative_error"] for r in ss],float)
      mm=np.nanmedian(errs); pass_rate=np.mean([r["pass_recovery"] for r in ss])
      wins=np.mean([r["relative_error"] < br["relative_error"] for r,br in zip(ss,[x for x in syn if x["model"]=="B"])]) if m!="B" else np.nan
      oo=[r for r in obs if r["model"]==m and np.isfinite(r["q_um"])]
      summary[m]={"synthetic_median_relative_error":float(mm),"synthetic_pass_rate":float(pass_rate),"win_rate_vs_B":None if m=="B" else float(wins),"observed_ok_runs":len(oo),"observed_total_runs":len([r for r in obs if r["model"]==m]),"evidence_status":"baseline" if m=="B" else ("supported" if mm<=.8*bmed and wins>=.75 else "inconclusive" if mm<bmed else "unsupported")}
    return summary

def json_safe(v):
    if isinstance(v, dict): return {k: json_safe(x) for k,x in v.items()}
    if isinstance(v, list): return [json_safe(x) for x in v]
    if isinstance(v, (np.bool_,)): return bool(v)
    if isinstance(v, (np.integer,)): return int(v)
    if isinstance(v, (np.floating, float)):
        return None if not np.isfinite(v) else float(v)
    return v

def main():
    started=time.perf_counter(); rows=synthetic_rows()+observed_rows(); summary=summarize(rows)
    fields=list(rows[0]);
    with (OUT/"原型结果.csv").open("w",newline="",encoding="utf-8-sig") as f:
      w=csv.DictWriter(f,fieldnames=fields); w.writeheader(); w.writerows(rows)
    for m in "BFP":
      d=OUT/f"model-{m.lower()}"; d.mkdir(parents=True,exist_ok=True)
      payload=json_safe({"model":m,"summary":summary[m],"runs":[r for r in rows if r["model"]==m]})
      (d/"output.json").write_text(json.dumps(payload,ensure_ascii=False,indent=2,allow_nan=False),encoding="utf-8")
    md=OUT/"model-m"; md.mkdir(parents=True,exist_ok=True)
    mdiag={"model":"M","role":"diagnostic_only","degenerate_limits":{"zero_back_reflection":"PASS_by_construction","strong_attenuation":"PASS_by_geometric_term_limit","incoherent_limit":"PASS_phase_term_removed"},"mechanism_gate":"INSUFFICIENT_EVIDENCE","thickness_effect_gate":"NOT_ESTIMABLE_WITHOUT_APPROVED_n_k","status":"GATE_NOT_TRIGGERED","admitted":False}
    (md/"output.json").write_text(json.dumps(mdiag,ensure_ascii=False,indent=2),encoding="utf-8")
    env={"python":sys.version,"platform":platform.platform(),"numpy":np.__version__,"openpyxl_module":"available","seed":SEED,"command":f'"{sys.executable}" "{Path(__file__).resolve()}"',"runtime_s":time.perf_counter()-started,"stop_rule":CFG["comparison"]["stopping_rule"]}
    (OUT/"environment.json").write_text(json.dumps(env,ensure_ascii=False,indent=2),encoding="utf-8")
    print(json.dumps({"summary":summary,"M":mdiag,"runtime_s":env["runtime_s"]},ensure_ascii=False,indent=2))

if __name__=="__main__": main()
